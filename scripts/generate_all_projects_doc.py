import os
import json
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
ALL_PROJ_DIR = os.path.dirname(SCRIPT_DIR)
CATALOG_PATH = os.path.join(ALL_PROJ_DIR, "projects_catalog.json")

with open(CATALOG_PATH, "r", encoding="utf-8") as f:
    catalog = json.load(f)

doc = docx.Document()

# Adjust margins
for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

# Title
title_p = doc.add_paragraph()
title_run = title_p.add_run("PORTFOLIO GLOBAL DES PROJETS D'INGÉNIERIE & IA")
title_run.font.name = "Calibri"
title_run.font.size = Pt(22)
title_run.font.bold = True
title_run.font.color.rgb = RGBColor(16, 44, 87)
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

sub_p = doc.add_paragraph()
sub_run = sub_p.add_run("Mohamed Nourou Dine CISSE — Élève-Ingénieur ENSIIE & Université Paris-Saclay (Master M2IA) | Échange McMaster University\nSpécialisation : Intelligence Artificielle, Deep Learning & Mathématiques Appliquées\nEmail : mouhamed.nouroudine.cisse@gmail.com | Tél : +33 7 65 60 09 97 | GitHub : github.com/nourou-dine-cisse")
sub_run.font.name = "Calibri"
sub_run.font.size = Pt(10.5)
sub_run.font.italic = True
sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph("-" * 80)

# Intro
p_intro = doc.add_paragraph()
r_intro = p_intro.add_run(f"Ce document rassemble la documentation exhaustive, les rapports d'expertise et les présentations de soutenance des {len(catalog)} projets d'ingénierie et de recherche réalisés au cours de mon cursus à l'ENSIIE, mon Master 2 en Intelligence Artificielle (Université Paris-Saclay), mon échange académique à McMaster University (Canada) et mes expériences industrielles (Publicis Groupe, Gainde 2000 / Infinity Analytics, Challenge Entreprendre).")
r_intro.font.name = "Calibri"
r_intro.font.size = Pt(11)

doc.add_heading(f"Tableau Récapitulatif des {len(catalog)} Projets", level=1)

table = doc.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr_cells = table.rows[0].cells
hdr_cells[0].text = "Projet"
hdr_cells[1].text = "Domaine / Institution"
hdr_cells[2].text = "Stack Technique"
hdr_cells[3].text = "Livrables (Rapports & Slides)"

for cell in hdr_cells:
    for p in cell.paragraphs:
        for r in p.runs:
            r.font.bold = True
            r.font.size = Pt(9.5)

for p in catalog:
    row_cells = table.add_row().cells
    row_cells[0].text = p["title_fr"].split(":")[0].strip()
    row_cells[1].text = f"{p['category']}\n({p['institution']})"
    row_cells[2].text = ", ".join(p["technologies"][:4])
    
    livrables = []
    if p["reports"]:
        livrables.append(f"📄 Rapport ({len(p['reports'])})")
    if p["presentations"]:
        livrables.append(f"📊 Présentation ({len(p['presentations'])})")
    if not livrables:
        livrables.append("💻 Code & Dépôt")
    row_cells[3].text = " + ".join(livrables)

    for cell in row_cells:
        for para in cell.paragraphs:
            for r in para.runs:
                r.font.size = Pt(9)

doc.add_page_break()

# Detailed Sections
for i, p in enumerate(catalog, 1):
    doc.add_heading(f"{i}. {p['title_fr']}", level=1)
    
    meta_p = doc.add_paragraph()
    m_run = meta_p.add_run(f"Contexte : {p['institution']} ({p['year']}) | Rôle : {p['role']}\nDomaine : {p['category']}\nTechnologies : {', '.join(p['technologies'])}")
    m_run.font.italic = True
    m_run.font.size = Pt(10)
    
    doc.add_heading("Synthèse Exécutive :", level=2)
    p_pitch = doc.add_paragraph()
    p_pitch.add_run(p["pitch_fr"])
    
    doc.add_heading("Réalisations Clés & Résultats (Format STAR) :", level=2)
    for b in p["star_bullet_points_fr"]:
        bp = doc.add_paragraph(style='List Bullet')
        bp.add_run(b)
        
    doc.add_heading("Livrables & Justificatifs Associés :", level=2)
    if p["reports"]:
        for r in p["reports"]:
            rp = doc.add_paragraph(style='List Bullet')
            rp.add_run(f"Rapport technique disponible : {r}")
    if p["presentations"]:
        for pr in p["presentations"]:
            prp = doc.add_paragraph(style='List Bullet')
            prp.add_run(f"Présentation / Soutenance : {pr}")
    if not p["reports"] and not p["presentations"]:
        np = doc.add_paragraph(style='List Bullet')
        np.add_run("Code source et documentation intégrés dans le dossier projet.")

    doc.add_paragraph("")

doc_path_all = os.path.join(ALL_PROJ_DIR, "All_Projects.docx")
doc_path_root = os.path.abspath("/Users/nouroudine/Documents/ENSIIE/All projects.docx")

doc.save(doc_path_all)
doc.save(doc_path_root)
print(f"[OK] Fichier Word généré dans : {doc_path_all}")
print(f"[OK] Fichier Word généré dans : {doc_path_root}")
